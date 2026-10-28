# -*- coding: utf-8 -*-
"""
Prioridades de pago que digitan los analistas
---------------------------------------------
Los analistas indican a que afiliados hay que priorizar el pago y por que, en un Excel compartido
("Prioridades de pago.xlsx", hoja "Prioridades"). actualizar_seguimiento.py lo lee en cada corrida y el
dashboard lo muestra (A quien pagar primero, Flujo de caja para Carlos, Prioridades de analistas).

El libro tiene tres hojas: Instrucciones, Prioridades (lo que se digita) y Listas (opciones de los desplegables).
Si el archivo no existe, se crea con los afiliados que mas importan ya listados por cliente.
"""
import difflib
import os
import re
import unicodedata
from datetime import datetime

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HOJA = "Prioridades"
OPC_PRIORIDAD = ["1 · Alta (pagar primero)", "2 · Media", "3 · Baja", "X · No pagar por ahora"]
OPC_MOTIVO = ["Acuerdo: si le pagan, carga", "Operación crítica (Ajover: transferencias, granel, entrega)", "Riesgo de perder al afiliado",
              "Cliente estratégico", "Deuda negociada con el afiliado", "Urgencia del afiliado", "Otro"]
OPC_COMPROMISO = ["Sí, carga si le pagan", "En negociación", "Sin compromiso"]
# (titulo, ancho, tipo): "info" = lo llena el sistema (gris); "entrada" = lo digita el analista (amarillo)
COLUMNAS = [("Cliente", 30, "info"), ("Afiliado", 40, "base"), ("Operación (Ajover)", 22, "info"), ("Prioridad del sistema", 18, "info"),
            ("Se le debe vencido", 17, "info"), ("Venta 30 días", 17, "info"),
            ("Prioridad del analista", 25, "entrada"), ("Motivo", 36, "entrada"), ("Observaciones", 55, "entrada"),
            ("¿Va a cargar si le pagan?", 24, "entrada"), ("Fecha del compromiso", 17, "entrada"), ("Monto a pagar solicitado", 19, "entrada"),
            ("Analista", 20, "entrada"), ("Fecha de registro", 17, "entrada")]
NIVEL_TXT = {1: "1 · Crítico", 2: "2 · Importante", 3: "3 · A programar", 4: "4 · Al ritmo", 5: "5 · Saldo antiguo", 0: "0 · Sin saldo"}
FILAS_VALIDACION = 3000


def norm(s):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode().upper()
    return re.sub(r"\s+", " ", s).strip()


def _nivel(n):
    try:
        return int(str(n)[0])
    except (ValueError, IndexError):
        return 0


# ----------------------------------------------------------------------------- CREAR EL LIBRO
def crear_plantilla(ruta, filas, nombres, hoy):
    """filas: DataFrame (cliente, afiliado, segmento, nivel, venc_libre, venta30). nombres: lista de todos los afiliados."""
    wb = Workbook()
    wi = wb.active
    wi.title = "Instrucciones"
    wp = wb.create_sheet(HOJA)
    wl = wb.create_sheet("Listas")

    # --- Listas (desplegables)
    wl.append(["Afiliados", "", "Prioridad", "Motivo", "Compromiso", "Operación (Ajover)"])
    ops = ["TRANSFERENCIAS", "GRANELES", "ENTREGA A CLIENTES", "TRANSFERENCIA CALI", "ZORROS CALI", "FIJOS CALI", "FIJOS MEDELLIN",
           "SENCILLOS MADRID", "TRANSFERENCIA ALAMO", "OTROS"]
    for i in range(max(len(nombres), 10) + 1):
        wl.cell(i + 2, 1, nombres[i] if i < len(nombres) else None)
        if i < len(OPC_PRIORIDAD):
            wl.cell(i + 2, 3, OPC_PRIORIDAD[i])
        if i < len(OPC_MOTIVO):
            wl.cell(i + 2, 4, OPC_MOTIVO[i])
        if i < len(OPC_COMPROMISO):
            wl.cell(i + 2, 5, OPC_COMPROMISO[i])
        if i < len(ops):
            wl.cell(i + 2, 6, ops[i])
    for c, w in zip("ABCDEF", [44, 3, 26, 52, 24, 24]):
        wl.column_dimensions[c].width = w
    for c in wl[1]:
        c.font = Font(bold=True)
    wl.sheet_state = "hidden"

    # --- Prioridades
    borde = Side(style="thin", color="D0D5DD")
    for j, (tit, ancho, tipo) in enumerate(COLUMNAS, 1):
        c = wp.cell(1, j, tit)
        c.font = Font(bold=True, color="FFFFFF" if tipo != "entrada" else "1B2433")
        c.fill = PatternFill("solid", fgColor="475467" if tipo != "entrada" else "FFD966")
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        wp.column_dimensions[get_column_letter(j)].width = ancho
    wp.row_dimensions[1].height = 34
    n = len(filas)
    for i, r in enumerate(filas.itertuples(index=False), 2):
        valores = [r.cliente, r.afiliado, r.segmento, NIVEL_TXT.get(_nivel(r.nivel), ""), round(float(r.venc_libre)), round(float(r.venta30))]
        for j, v in enumerate(valores, 1):
            c = wp.cell(i, j, v)
            if j in (5, 6):
                c.number_format = "#,##0"
            if COLUMNAS[j - 1][2] == "info":
                c.fill = PatternFill("solid", fgColor="F2F4F7")
                c.font = Font(color="475467")
    for i in range(2, FILAS_VALIDACION + 1):
        for j, (tit, ancho, tipo) in enumerate(COLUMNAS, 1):
            c = wp.cell(i, j)
            if tipo == "entrada":
                c.fill = PatternFill("solid", fgColor="FFFBEA")
            c.border = Border(bottom=borde)
            if j == 9:
                c.alignment = Alignment(wrap_text=True, vertical="top")
            if j in (11, 14):
                c.number_format = "dd/mm/yyyy"
            if j == 12:
                c.number_format = "#,##0"
    nl = max(len(nombres), 1) + 1
    dvs = [("B", f"=Listas!$A$2:$A${nl}"), ("G", "=Listas!$C$2:$C$5"), ("H", f"=Listas!$D$2:$D${len(OPC_MOTIVO) + 1}"),
           ("J", f"=Listas!$E$2:$E${len(OPC_COMPROMISO) + 1}")]
    for col, formula in dvs:
        dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=(col != "B"))
        if col == "B":
            dv.showErrorMessage = False        # se permite escribir un afiliado que no este en la lista
        dv.error = "Elige una opción de la lista"
        wp.add_data_validation(dv)
        dv.add(f"{col}2:{col}{FILAS_VALIDACION}")
    rango = f"G2:G{FILAS_VALIDACION}"
    for letra, color in [("1", "FFD9D2"), ("2", "FFF1C2"), ("3", "E3ECF7")]:
        wp.conditional_formatting.add(rango, FormulaRule(formula=[f'LEFT($G2,1)="{letra}"'], fill=PatternFill("solid", bgColor=color, fgColor=color)))
    wp.conditional_formatting.add(rango, FormulaRule(formula=['UPPER(LEFT($G2,1))="X"'], fill=PatternFill("solid", bgColor="D9D9D9", fgColor="D9D9D9")))
    wp.freeze_panes = "C2"
    wp.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNAS))}{FILAS_VALIDACION}"

    # --- Instrucciones
    wi.column_dimensions["A"].width = 120
    lineas = [
        ("PRIORIDADES DE PAGO A AFILIADOS · Tractocar", "titulo"),
        ("", None),
        ("Para qué sirve", "sub"),
        ("Aquí los analistas indican a qué afiliados hay que priorizar el pago y por qué. El dashboard (Centro de control) lee la hoja «Prioridades» cada vez que se actualiza: "
         "el plan «A quién pagar primero» las respeta, el «Flujo de caja» se las muestra a Carlos y la pestaña «Prioridades de analistas» las lista por cliente.", None),
        ("", None),
        ("Cómo diligenciar", "sub"),
        ("1. En la hoja «Prioridades» filtra por tu cliente (columna A). Las columnas grises son informativas (corresponden a la fecha de creación del archivo); no las modifiques.", None),
        ("2. En la fila del afiliado elige la «Prioridad del analista»: 1 Alta (pagar primero), 2 Media, 3 Baja o X No pagar por ahora.", None),
        ("3. Elige el «Motivo» y escribe en «Observaciones» por qué. Ejemplo: «Se acordó con el afiliado que si se le paga el saldo, carga 10 viajes esta semana».", None),
        ("4. Si hay un acuerdo de carga, elige «Sí, carga si le pagan» y la «Fecha del compromiso».", None),
        ("5. Si quieres un valor específico, escribe el «Monto a pagar solicitado» (en pesos). Si lo dejas vacío, el plan decide cuánto.", None),
        ("6. Escribe tu nombre en «Analista» y la «Fecha de registro».", None),
        ("", None),
        ("Reglas", "sub"),
        ("• Una sola fila por afiliado. Si hay dos, el dashboard usa la de «Fecha de registro» más reciente.", None),
        ("• Solo cuentan las filas que tienen «Prioridad del analista». «Alta» sube al afiliado al inicio del plan (aunque el sistema no lo priorice); «Baja» lo deja al final; «No pagar por ahora» lo saca del plan.", None),
        ("• Si un afiliado no está en la lista, escribe en una fila nueva (al final) el nombre tal como aparece en el sistema; el dashboard avisa si no lo encuentra.", None),
        ("• Para ver el efecto, actualiza el dashboard (ACTUALIZAR.bat).", None),
        ("", None),
        (f"Archivo creado el {hoy:%d/%m/%Y}. Información del sistema (columnas grises) a esa fecha.", None),
    ]
    for i, (t, estilo) in enumerate(lineas, 1):
        c = wi.cell(i, 1, t)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if estilo == "titulo":
            c.font = Font(bold=True, size=16, color="1F3864")
        elif estilo == "sub":
            c.font = Font(bold=True, size=12, color="1F3864")
    wb.save(ruta)
    return ruta


def filas_plantilla(cja, ccj, seg_af, min_antiguo=2e6):
    """Afiliados que se listan por cliente: los que cargan (niveles 1-4) y el saldo antiguo de al menos min_antiguo."""
    nivel = cja["nivel"].map(_nivel)
    sel = cja[(nivel.between(1, 4)) | ((nivel == 5) & (cja["venc_libre"] >= min_antiguo))].copy()
    cli = (ccj.sort_values("venta", ascending=False).drop_duplicates("afiliado").set_index("afiliado")["cliente"]) if len(ccj) else pd.Series(dtype=object)
    sel["cliente"] = sel["afiliado"].map(cli).fillna("(sin viajes en 30 días)")
    sel["segmento"] = sel["afiliado"].map(lambda a: seg_af.get(norm(a), ""))
    sel["_sin"] = (sel["cliente"] == "(sin viajes en 30 días)").astype(int)          # los que no cargan, al final
    sel = sel.sort_values(["_sin", "cliente", "venc_libre"], ascending=[True, True, False])
    return sel[["cliente", "afiliado", "segmento", "nivel", "venc_libre", "venta30"]]


# ----------------------------------------------------------------------------- LEER LAS PRIORIDADES
def _prio(v):
    t = str(v).strip().upper()
    if not t or t == "NAN":
        return None
    return {"1": 1, "2": 2, "3": 3, "X": 0, "0": 0, "N": 0}.get(t[0])


def _fecha(v):
    d = pd.to_datetime(v, errors="coerce", dayfirst=True)
    return None if pd.isna(d) else d.strftime("%Y-%m-%d")


def _emparejar(texto, canon):
    """canon: {afn: nombre}. Devuelve (nombre, estado) con estado OK / APROXIMADO / AMBIGUO / NO ENCONTRADO."""
    k = norm(texto)
    if not k:
        return None, "NO ENCONTRADO"
    if k in canon:
        return canon[k], "OK"
    cont = [a for a in canon if k in a or a in k]
    if len(cont) == 1:
        return canon[cont[0]], "APROXIMADO"
    toks = set(k.split())
    sub = [a for a in canon if toks <= set(a.split())]
    if len(sub) == 1:
        return canon[sub[0]], "APROXIMADO"
    if len(cont) > 1 or len(sub) > 1:
        return None, "AMBIGUO"
    cerca = difflib.get_close_matches(k, list(canon), n=1, cutoff=0.88)
    if cerca:
        return canon[cerca[0]], "APROXIMADO"
    return None, "NO ENCONTRADO"


def leer(ruta, canon):
    """Lee la hoja Prioridades. Devuelve (lista de dicts, meta)."""
    meta = dict(archivo=ruta, existe=os.path.exists(ruta), modificado="", n=0, alta=0, media=0, baja=0, no_pagar=0, sin_match=0, sin_prioridad=0, duplicados=0)
    if not meta["existe"]:
        return [], meta
    meta["modificado"] = datetime.fromtimestamp(os.path.getmtime(ruta)).strftime("%d/%m/%Y %H:%M")
    try:
        x = pd.read_excel(ruta, sheet_name=HOJA, dtype=object)
    except Exception as ex:                      # archivo abierto/bloqueado o sin la hoja
        meta["error"] = f"{type(ex).__name__}: {ex}"
        return [], meta
    cols = {norm(c): c for c in x.columns}

    def col(nombre):
        return x[cols[norm(nombre)]] if norm(nombre) in cols else pd.Series([None] * len(x), index=x.index)

    d = pd.DataFrame({"cliente": col("Cliente"), "afiliado_dig": col("Afiliado"), "prio_txt": col("Prioridad del analista"), "motivo": col("Motivo"),
                      "obs": col("Observaciones"), "compromiso": col("¿Va a cargar si le pagan?"), "f_comp": col("Fecha del compromiso"),
                      "monto": col("Monto a pagar solicitado"), "analista": col("Analista"), "f_reg": col("Fecha de registro")})
    d["fila"] = d.index + 2
    d = d[d["afiliado_dig"].notna() & (d["afiliado_dig"].astype(str).str.strip() != "")]
    d["prioridad"] = d["prio_txt"].map(_prio)
    meta["sin_prioridad"] = int(d["prioridad"].isna().sum())
    d = d[d["prioridad"].notna()].copy()
    em = d["afiliado_dig"].map(lambda t: _emparejar(t, canon))
    d["afiliado"] = [a if a else str(t).strip() for (a, _), t in zip(em, d["afiliado_dig"])]
    d["estado_match"] = [e for _, e in em]
    d["f_reg_iso"] = d["f_reg"].map(_fecha)
    d["f_comp_iso"] = d["f_comp"].map(_fecha)
    d["monto_n"] = pd.to_numeric(d["monto"], errors="coerce").fillna(0)
    d = d.sort_values(["f_reg_iso", "fila"], na_position="first")
    ya = d["estado_match"] != "NO ENCONTRADO"
    dup = d[ya].duplicated("afiliado", keep="last")
    meta["duplicados"] = int(dup.sum())
    d = d.drop(index=d[ya][dup].index)
    txt = lambda v: "" if v is None or (isinstance(v, float) and np.isnan(v)) else str(v).strip()
    filas = [dict(afiliado=r.afiliado, afiliado_digitado=txt(r.afiliado_dig), estado_match=r.estado_match, cliente=txt(r.cliente), prioridad=int(r.prioridad),
                  prioridad_txt=txt(r.prio_txt), motivo=txt(r.motivo), obs=txt(r.obs), compromiso=txt(r.compromiso), fecha_compromiso=r.f_comp_iso or "",
                  monto=float(r.monto_n), analista=txt(r.analista), fecha=r.f_reg_iso or "", fila=int(r.fila)) for r in d.itertuples(index=False)]
    meta.update(n=len(filas), alta=sum(f["prioridad"] == 1 for f in filas), media=sum(f["prioridad"] == 2 for f in filas), baja=sum(f["prioridad"] == 3 for f in filas),
                no_pagar=sum(f["prioridad"] == 0 for f in filas), sin_match=sum(f["estado_match"] == "NO ENCONTRADO" for f in filas))
    return filas, meta


def preparar(ruta, cja, ccj, seg_af, hoy, log=print):
    """Crea el libro si no existe (con los afiliados listados por cliente) y lo lee. cja: afiliado, afn, nivel, venc_libre, venta30."""
    creado = False
    if not os.path.exists(ruta):
        try:
            crear_plantilla(ruta, filas_plantilla(cja, ccj, seg_af), sorted(cja["afiliado"].astype(str).unique()), hoy)
            creado = True
            log(f"Se creo el Excel de prioridades de los analistas: {ruta}")
        except Exception as ex:
            log(f"No se pudo crear el Excel de prioridades ({type(ex).__name__}: {ex})")
    canon = {norm(a): a for a in cja["afiliado"].astype(str).unique()}
    filas, meta = leer(ruta, canon)
    meta["creado_ahora"] = creado
    return filas, meta
